"""Write result workbooks and v11 diagnostic files.

Template column order is kept; slot k fills position k. Template headers are
not rewritten.
"""

from __future__ import annotations

import json
from datetime import datetime, timedelta
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import Workbook, load_workbook

from .constants import EXPORT_DAYS, EXPORT_START, N_DAYS, T, TEMPLATE_DIR
from .data import Data, time_map_rows
from .simulate import YearRun


BLOCKS = [
    (0, 24, "0:00-4:00"),
    (24, 48, "4:00-8:00"),
    (48, 72, "8:00-12:00"),
    (72, 96, "12:00-16:00"),
    (96, 120, "16:00-20:00"),
    (120, 144, "20:00-24:00"),
]


def _dates():
    d0 = datetime(2025, 1, 1)
    return [d0 + timedelta(days=i) for i in range(365)]


def _fmt_clock(minutes: int) -> str:
    if minutes >= 1440:
        return "24:00"
    return f"{minutes // 60}:{minutes % 60:02d}"


def merge_emergency(h: np.ndarray, tol: float = 1e-6) -> list[tuple[str, float]]:
    out = []
    t = 0
    n = len(h)
    while t < n:
        if h[t] > tol:
            t0 = t
            s = 0.0
            while t < n and h[t] > tol:
                s += float(h[t])
                t += 1
            out.append((f"{_fmt_clock(t0 * 10)}-{_fmt_clock(t * 10)}", s))
        else:
            t += 1
    return out


def export_result1(path, labels: list[str], a, c, d, E0, E24) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "计划购电量"
    ws.append(["时间段", "购电量"])
    for lab, val in zip(labels, a):
        ws.append([lab, float(val)])
    ws2 = wb.create_sheet("充放电量")
    ws2.append(["时间段", "充电量", "放电量", "时刻", "储电量"])
    for i, (lo, hi, name) in enumerate(BLOCKS):
        row = [name, float(c[lo:hi].sum()), float(d[lo:hi].sum()), None, None]
        if i == 0:
            row[3], row[4] = "0:00", float(E0)
        if i == 1:
            row[3], row[4] = "24:00", float(E24)
        ws2.append(row)
    wb.save(path)


def _write_plan_sheet(ws, headers, run: YearRun, settle_cost, settle_qty):
    ws.append(headers)
    dates = _dates()
    n_export = min(EXPORT_DAYS, max(0, run.q.shape[0] - EXPORT_START))
    for d in range(EXPORT_START, EXPORT_START + n_export):
        row = [dates[d]] + [float(x) for x in run.q[d]] + [float(settle_qty[d]), float(settle_cost[d])]
        if len(row) < len(headers):
            row.extend([None] * (len(headers) - len(row)))
        else:
            row = row[: len(headers)]
        ws.append(row)


def _write_adjust_sheet(ws, headers, run: YearRun):
    ws.append(headers)
    dates = _dates()
    n_export = min(EXPORT_DAYS, max(0, run.a.shape[0] - EXPORT_START))
    for d in range(EXPORT_START, EXPORT_START + n_export):
        qty = float(run.a[d].sum())
        cost = float(run.day_cost[d])
        row = [dates[d]] + [float(x) for x in run.a[d]] + [qty, cost]
        if len(row) < len(headers):
            row.extend([None] * (len(headers) - len(row)))
        else:
            row = row[: len(headers)]
        ws.append(row)


def _write_bess_sheet(ws, run: YearRun):
    ws.append(["日期", "时间段", "充电量", "放电量", "时刻", "储电量"])
    dates = _dates()
    n_export = min(EXPORT_DAYS, max(0, run.q.shape[0] - EXPORT_START))
    for d in range(EXPORT_START, EXPORT_START + n_export):
        for i, (lo, hi, name) in enumerate(BLOCKS):
            row = [
                dates[d] if i == 0 else None,
                name,
                float(run.c[d, lo:hi].sum()),
                float(run.d[d, lo:hi].sum()),
                None,
                None,
            ]
            if i == 0:
                row[4], row[5] = "0:00", float(run.E_start[d])
            if i == 1:
                row[4], row[5] = "24:00", float(run.E_end[d])
            ws.append(row)


def _write_emerg_sheet(ws, run: YearRun):
    ws.append(["日期", "购电时间段", "购电量"])
    dates = _dates()
    n_export = min(EXPORT_DAYS, max(0, run.q.shape[0] - EXPORT_START))
    for d in range(EXPORT_START, EXPORT_START + n_export):
        ivs = merge_emergency(run.h[d])
        if not ivs:
            ws.append([dates[d], None, None])
            continue
        for j, (span, qty) in enumerate(ivs):
            ws.append([dates[d] if j == 0 else None, span, float(qty)])


def _headers_from_template(filename: str, sheet: str) -> list:
    wb = load_workbook(TEMPLATE_DIR / filename, read_only=True, data_only=True)
    ws = wb[sheet]
    headers = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
    wb.close()
    return headers


def export_year(path, template_name: str, run: YearRun, settle_cost, settle_qty, with_adjust: bool) -> None:
    headers_q = _headers_from_template(template_name, "计划购电量")
    wb = Workbook()
    ws = wb.active
    ws.title = "计划购电量"
    _write_plan_sheet(ws, headers_q, run, settle_cost, settle_qty)
    if with_adjust:
        headers_a = _headers_from_template(template_name, "调整购电量")
        ws_a = wb.create_sheet("调整购电量")
        _write_adjust_sheet(ws_a, headers_a, run)
    ws_b = wb.create_sheet("充放电量")
    _write_bess_sheet(ws_b, run)
    ws_h = wb.create_sheet("紧急购电量")
    _write_emerg_sheet(ws_h, run)
    wb.save(path)


def export_time_map(path: Path, data: Data) -> None:
    df = pd.DataFrame(time_map_rows(data))
    df.to_csv(path, index=False, encoding="utf-8")


def export_trace_npz(path: Path, run: YearRun, data: Data) -> None:
    np.savez_compressed(
        path,
        q=run.q,
        a=run.a,
        c=run.c,
        d=run.d,
        h=run.h,
        w=run.w,
        E_slot=run.E_slot,
        p_plan=run.p_plan,
        p_set=run.p_set,
        fc_L0=run.fc_L0,
        fc_V0=run.fc_V0,
        a_snap=run.a_snap,
        c_opt=run.c_opt,
        d_opt=run.d_opt,
        src=run.src,
        clip_reason=run.clip_reason,
        load=data.load[: run.q.shape[0]],
        pv=data.pv[: run.q.shape[0]],
        day_cost=run.day_cost,
    )


def export_plans_jsonl(path: Path, run: YearRun) -> None:
    dates = _dates()
    with path.open("w", encoding="utf-8") as f:
        for ev in run.plan_events:
            rec = dict(ev)
            d = int(rec["day"])
            rec["date"] = dates[d].strftime("%Y-%m-%d") if d < len(dates) else str(d)
            f.write(json.dumps(rec, ensure_ascii=False, default=str) + "\n")


def export_specified_csv(path: Path, run: YearRun, data: Data, dates_keep: tuple[str, ...]) -> None:
    dates = pd.to_datetime(data.dates[: run.q.shape[0]])
    rows = []
    keep = set(dates_keep)
    for d in range(run.q.shape[0]):
        date_s = pd.Timestamp(dates[d]).strftime("%Y-%m-%d")
        if date_s not in keep:
            continue
        for t in range(T):
            rows.append(
                {
                    "date": date_s,
                    "slot": t,
                    "interval": f"[{t * 10 // 60}:{t * 10 % 60:02d}, {(t + 1) * 10 // 60}:{(t + 1) * 10 % 60:02d})",
                    "q": float(run.q[d, t]),
                    "a": float(run.a[d, t]),
                    "h": float(run.h[d, t]),
                    "c": float(run.c[d, t]),
                    "dchg": float(run.d[d, t]),
                    "w": float(run.w[d, t]),
                    "E_start": float(run.E_slot[d, t]),
                    "E_end": float(run.E_slot[d, t + 1]),
                    "load_kw": float(data.load[d, t]),
                    "pv_kw": float(data.pv[d, t]),
                    "fc_L0_kw": float(run.fc_L0[d, t]),
                    "fc_V0_kw": float(run.fc_V0[d, t]),
                    "p_plan": float(run.p_plan[d, t]),
                    "p_set": float(run.p_set[d, t]),
                    "c_opt": float(run.c_opt[d, t]),
                    "d_opt": float(run.d_opt[d, t]),
                    "src": int(run.src[d, t]),
                    "clip_reason": int(run.clip_reason[d, t]),
                }
            )
    pd.DataFrame(rows).to_csv(path, index=False, encoding="utf-8")


def export_snapshots_jsonl(path: Path, run: YearRun) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8") as f:
        for rec in run.snapshots:
            f.write(json.dumps(rec, ensure_ascii=False, default=str) + "\n")


def export_audit(path: Path, summaries: list[dict]) -> None:
    path.write_text(json.dumps(summaries, indent=2, ensure_ascii=False), encoding="utf-8")
